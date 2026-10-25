import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
import cashflow  # noqa: E402
import make_example  # noqa: E402


@pytest.fixture
def book(tmp_path, monkeypatch):
    monkeypatch.setenv("OPS_DESK_DATA", str(tmp_path / "ops"))
    path = tmp_path / "cashflow-example.xlsx"
    make_example.build(path)
    return path


def run(capsys, book, *args):
    code = cashflow.main(["--file", str(book), *args])
    return code, json.loads(capsys.readouterr().out)


def cells(book, sheet="Cash Flow"):
    return openpyxl.load_workbook(book)[sheet]


def add(capsys, book, when, desc, amt, kind="payment", head="Project Alpha", *extra):
    return run(capsys, book, "add", "--date", when, "--desc", desc, "--head", head,
               "--amount", str(amt), "--kind", kind, *extra)


def test_check_passes_on_example(book, capsys):
    code, out = run(capsys, book, "check")
    assert code == 0 and out["ok"]


def test_summary_matches_hand_calculation(book, capsys):
    _, out = run(capsys, book, "summary")
    assert (out["opening_inr"], out["closing_inr"], out["lowest_inr"], out["lowest_on"]) == (1350000, 160000, 160000, "2027-03-31")
    oct_ = out["months"][0]
    assert (oct_["receipts_inr"], oct_["payments_inr"], oct_["closing_inr"], oct_["lowest_inr"]) == (200000, 100000, 1450000, 1250000)
    assert out["pending_by_head"]["Monthly"]["payments_inr"] == 1200000


def test_add_inserts_in_date_order_and_updates_ranges(book, capsys):
    code, out = add(capsys, book, "2026-11-20", "Transport", 50000)
    assert code == 0 and out["rows"][0]["row"] == 11
    assert out["after"]["closing_inr"] == 110000 and out["saved"]
    ws = cells(book)
    assert ws["B11"].value == "Transport" and ws["E11"].value == "=E10-N(D11)"
    assert ws["E12"].value == "=E11-N(D12)"
    ms = cells(book, "Month Summary")
    assert "'Cash Flow'!$D$5:$D$14" in ms["C5"].value
    assert "'Cash Flow'!E15" in ms["G11"].value
    assert {str(r.sqref) for r in ws.conditional_formatting} == {"F6:F14", "A6:C14", "E5:E14"}
    assert ws.print_area.endswith("$G$16")
    assert run(capsys, book, "check")[1]["ok"]


def test_add_before_recurring_base_moves_links(book, capsys):
    add(capsys, book, "2026-10-30", "Early item", 10000)
    ws = cells(book)
    assert ws["B11"].value == "Salary Sep"
    assert ws["D12"].value == "=D11" and ws["D13"].value == "=D11"
    assert run(capsys, book, "check")[1]["ok"]


def test_same_date_payment_goes_before_receipt(book, capsys):
    _, out = add(capsys, book, "2026-10-26", "Extra payment", 1000)
    assert out["rows"][0]["row"] == 9
    _, out = add(capsys, book, "2026-10-26", "Extra receipt", 1000, "receipt", "Receipt")
    assert out["rows"][0]["row"] == 11


def test_add_at_end_extends_ranges(book, capsys):
    code, out = add(capsys, book, "2027-03-31", "Last item", 5000)
    assert code == 0 and out["rows"][0]["row"] == 14
    assert "'Cash Flow'!$D$5:$D$14" in cells(book, "Month Summary")["C5"].value
    assert out["after"]["closing_inr"] == 155000


def test_add_rejects_duplicate_and_unknown_head(book, capsys):
    code, out = add(capsys, book, "2026-10-26", "Project Alpha fittings", 100000)
    assert code == 1 and "already has this item" in out["errors"][0]
    code, out = add(capsys, book, "2026-11-20", "x", 10, "payment", "Typo Head")
    assert code == 1 and "unknown head" in out["errors"][0]
    code, _ = add(capsys, book, "2026-11-20", "x", 10, "payment", "New Head", "--allow-new-head")
    assert code == 0


def test_remove_needs_matching_text(book, capsys):
    code, out = run(capsys, book, "remove", "--item", "9=Salary")
    assert code == 1 and "not 'Salary'" in out["errors"][0]


def test_remove_recurring_base_is_refused(book, capsys):
    code, out = run(capsys, book, "remove", "--item", "10=Salary Sep")
    assert code == 1 and "base of recurring" in out["errors"][0]


def test_remove_last_row_shrinks_ranges(book, capsys):
    code, out = run(capsys, book, "remove", "--item", "13=Software")
    assert code == 0 and out["after"]["closing_inr"] == 250000
    assert "'Cash Flow'!$D$5:$D$12" in cells(book, "Month Summary")["C5"].value
    assert run(capsys, book, "check")[1]["ok"]


def test_amount_change_reports_followers(book, capsys):
    code, out = run(capsys, book, "amount", "--item", "10=Salary Sep", "--amount", "500000")
    assert code == 0 and out["rows_following_this_amount"] == [11, 12]
    assert out["after"]["closing_inr"] == -140000
    code, out = run(capsys, book, "amount", "--item", "11=Salary Oct", "--amount", "1")
    assert code == 1 and "follows another row" in out["errors"][0]


def test_move_in_place_only_within_date_order(book, capsys):
    code, _ = run(capsys, book, "move", "--item", "8=Project Alpha", "--date", "2026-10-25")
    assert code == 0
    code, out = run(capsys, book, "move", "--item", "8=Project Alpha", "--date", "2026-10-27")
    assert code == 1 and "outside" in out["errors"][0]


def test_mark_done_sets_opening_balance_and_date(book, capsys):
    code, out = run(capsys, book, "mark-done", "--item", "8=Project Alpha",
                    "--opening-balance", "1250000", "--as-of", "2026-10-26")
    assert code == 0 and out["after"]["closing_inr"] == 160000
    ws = cells(book)
    assert ws["D3"].value == 1250000 and ws["F8"].value == 1 and ws["E8"].value == "=E7"
    assert "Done on 26.10.2026" in ws["G8"].value
    assert cells(book, "Month Summary")["H5"].value.date().isoformat() == "2026-10-26"
    code, out = run(capsys, book, "mark-done", "--item", "10=Salary Sep",
                    "--opening-balance", "1", "--as-of", "2026-10-26")
    assert code == 1 and "dated after" in out["errors"][0]


def test_dry_run_leaves_file_untouched(book, capsys):
    before = book.read_bytes()
    code, out = add(capsys, book, "2026-11-20", "Transport", 50000, "payment", "Project Alpha", "--dry-run")
    assert code == 0 and not out["saved"] and book.read_bytes() == before


def test_save_makes_a_backup(book, capsys, tmp_path):
    _, out = add(capsys, book, "2026-11-20", "Transport", 50000)
    assert Path(out["backup"]).exists() and Path(out["backup"]).parent == tmp_path / "ops" / "backups"


def test_refuses_to_write_when_excel_has_file_open(book, capsys):
    (book.parent / "~$shflow-example.xlsx").write_text("lock")
    code, out = add(capsys, book, "2026-11-20", "Transport", 50000)
    assert code == 1 and "open in Excel" in out["errors"][0]


def test_broken_balance_formula_blocks_edits(book, capsys):
    wb = openpyxl.load_workbook(book)
    wb["Cash Flow"]["E9"] = "=E8"
    wb.save(book)
    code, out = add(capsys, book, "2026-11-20", "Transport", 50000)
    assert code == 1 and any("E9" in e for e in out["errors"])


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_libreoffice_recalculation_agrees_after_edits(book, capsys, tmp_path):
    add(capsys, book, "2026-10-30", "Early item", 10000)
    run(capsys, book, "remove", "--item", "14=Software")
    add(capsys, book, "2027-03-31", "Last item", 5000)
    out_dir = tmp_path / "recalc"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(book)],
                   check=True, capture_output=True, timeout=120)
    values = openpyxl.load_workbook(out_dir / book.name, data_only=True)
    ws, ms = values["Cash Flow"], values["Month Summary"]
    close = cashflow.close_row(cells(book))
    mine = cashflow.key_figures(cells(book))
    assert ws[f"E{close}"].value == pytest.approx(mine["closing_inr"])
    assert ws[f"E{close + 1}"].value == pytest.approx(mine["lowest_inr"])
    assert ms["G11"].value == "OK"
    months = cashflow.months(openpyxl.load_workbook(book), cells(book))
    assert [ms[f"E{r}"].value for r in range(5, 11)] == pytest.approx([m["closing_inr"] for m in months])
